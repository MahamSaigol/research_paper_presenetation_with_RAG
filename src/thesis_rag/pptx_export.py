"""
Builds a .pptx from SlideDraft content (slidegen.py) — this is the actual
deliverable, since a Python-drafted summary is only useful to you as a
real slide file you send to your supervisor, not as text in a terminal.

Design choice: assign text via `run.text`, never `text_frame.text = ...`.
    Overwriting a text frame's `.text` attribute directly collapses
    whatever paragraph/run structure exists into one unstyled run — you
    lose the ability to set per-bullet formatting (e.g. italicizing the
    "not stated" case) afterward. Building paragraphs and runs
    explicitly avoids that.

Design choice: each slide's source pages are printed in a small footer.
    This is a *draft* you're meant to correct — the footer is what lets
    you jump straight to the right page(s) in the source PDF to verify
    or expand what the model wrote, without hunting for it.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .slidegen import SlideDraft

_NOT_STATED_TEXT = "Not clearly stated in this paper."
_GRAY = RGBColor(0x80, 0x80, 0x80)


def build_pptx(paper_name: str, drafts: list[SlideDraft], output_path: Path) -> None:
    prs = Presentation()

    # Strip the file extension for the title slide — "43. WeedNet-X....pdf"
    # reads as a filename, not a paper title, in a deck meant for your
    # supervisor. Only the extension is stripped, not any numbering
    # prefix — filenames vary too much across your library (some have a
    # "43. " prefix, some don't) to safely guess at further cleanup.
    display_title = Path(paper_name).stem

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = display_title
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "Draft summary — review before use"

    content_layout = prs.slide_layouts[1]  # "Title and Content"

    for draft in drafts:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = draft.title

        body = slide.placeholders[1].text_frame
        body.clear()  # leaves exactly one empty paragraph to build on

        if draft.bullets:
            for i, bullet_text in enumerate(draft.bullets):
                paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
                run = paragraph.add_run()
                run.text = bullet_text
        else:
            run = body.paragraphs[0].add_run()
            run.text = _NOT_STATED_TEXT
            run.font.italic = True
            run.font.color.rgb = _GRAY

        if draft.source_pages:
            footer_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(6.8), Inches(9), Inches(0.4)
            )
            footer_run = footer_box.text_frame.paragraphs[0].add_run()
            footer_run.text = f"Source: {', '.join(draft.source_pages)}"
            footer_run.font.size = Pt(10)
            footer_run.font.color.rgb = _GRAY

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
